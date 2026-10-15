# -*- coding: utf-8 -*-
"""
Generador del documento Word para Runner Challenge
Portada + contenido académico para presentación al docente
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = r"c:\Users\Henry Camargo\Desktop\GameRunnig\Runner_Challenge_Presentacion.docx"

doc = Document()

# ── Márgenes del documento ───────────────────────────────────────────────────
from docx.oxml.ns import qn
section = doc.sections[0]
section.page_width  = Cm(21.59)   # Carta
section.page_height = Cm(27.94)
section.top_margin    = Cm(3)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3)
section.right_margin  = Cm(2.5)

# ── Colores del tema ──────────────────────────────────────────────────────────
AZUL_OSCURO  = RGBColor(0x0D, 0x1B, 0x2A)   # Fondo oscuro
CYAN_NEON    = RGBColor(0x00, 0xB4, 0xD8)   # Cian del juego
GRIS_CLARO   = RGBColor(0xF0, 0xF0, 0xF0)
NEGRO        = RGBColor(0x00, 0x00, 0x00)
AZUL_MED     = RGBColor(0x02, 0x3E, 0x8A)

def set_cell_bg(cell, hex_color):
    """Pone color de fondo a una celda de tabla."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def add_paragraph_border_bottom(para):
    """Agrega línea decorativa debajo de un párrafo."""
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '00B4D8')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ════════════════════════════════════════════════════════════════════════════
#  PÁGINA 1 — PORTADA
# ════════════════════════════════════════════════════════════════════════════

# Espacio inicial
for _ in range(2):
    doc.add_paragraph()

# ── Título principal del proyecto ────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("RUNNER CHALLENGE")
run.bold      = True
run.font.size = Pt(36)
run.font.color.rgb = AZUL_MED

# Subtítulo
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Neon Run — Videojuego 2D Runner")
r2.font.size = Pt(18)
r2.font.color.rgb = CYAN_NEON
r2.bold = True
add_paragraph_border_bottom(p2)

doc.add_paragraph()

# ── Etiqueta del tipo de documento ───────────────────────────────────────────
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run("Avance 1 — Prototipo Funcional")
r3.font.size = Pt(14)
r3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
r3.italic = True

for _ in range(3):
    doc.add_paragraph()

# ── Tabla de datos de presentación ───────────────────────────────────────────
table = doc.add_table(rows=6, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'

filas = [
    ("Asignatura",   "Electiva de Profundizacion"),
    ("Docente",      "Cristian Torres"),
    ("Estudiantes",  "Brandon Lara"),
    ("",             "Henry Camargo"),
    ("Semestre",     "2026 - II"),
    ("Entrega",      "Septiembre 2026"),
]

for i, (label, valor) in enumerate(filas):
    row = table.rows[i]
    # Columna etiqueta
    cell_label = row.cells[0]
    set_cell_bg(cell_label, "023E8A")
    p_label = cell_label.paragraphs[0]
    p_label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_label = p_label.add_run(label)
    run_label.bold = True
    run_label.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run_label.font.size = Pt(12)

    # Columna valor
    cell_val = row.cells[1]
    set_cell_bg(cell_val, "F0F4FF")
    p_val = cell_val.paragraphs[0]
    p_val.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_val = p_val.add_run(valor)
    run_val.font.size = Pt(12)
    run_val.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    if label == "Docente" or label == "Asignatura":
        run_val.bold = True

# Ancho de columnas
for row in table.rows:
    row.cells[0].width = Cm(5)
    row.cells[1].width = Cm(9)

for _ in range(4):
    doc.add_paragraph()

# ── Pie de portada ────────────────────────────────────────────────────────────
p_foot = doc.add_paragraph()
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_foot = p_foot.add_run("Proyecto de Desarrollo de Videojuego — Programa de Ingenieria")
r_foot.font.size = Pt(10)
r_foot.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
r_foot.italic = True

# Salto de página antes del contenido
doc.add_page_break()


# ════════════════════════════════════════════════════════════════════════════
#  PÁGINA 2 — DESCRIPCIÓN DEL PROYECTO
# ════════════════════════════════════════════════════════════════════════════

def heading(doc, text, level=1, color=AZUL_MED):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = color
    return p

def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.size = Pt(11)
    return p

heading(doc, "1. Descripcion del Proyecto", 1)
body(doc,
    "Runner Challenge: Neon Run es un videojuego 2D de tipo runner desarrollado como "
    "prototipo funcional (Avance 1) para la asignatura Electiva de Profundizacion. "
    "El juego consiste en controlar uno o dos personajes que corren automaticamente "
    "por un escenario con estetica cyberpunk, evitando obstaculos y recolectando "
    "monedas para acumular puntos."
)

doc.add_paragraph()
heading(doc, "2. Objetivos", 1)
objetivos = [
    "Demostrar la implementacion de mecanicas basicas de un videojuego 2D (salto, colision, puntuacion).",
    "Aplicar conceptos de programacion orientada a objetos (clases Player, ObstacleSystem, Camera).",
    "Implementar un bucle de juego (game loop) con fisicas simples de gravedad.",
    "Desarrollar un modo multijugador local en el mismo teclado.",
    "Presentar un prototipo funcional como avance del proyecto final.",
]
for obj in objetivos:
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(obj)
    run.font.size = Pt(11)

doc.add_paragraph()
heading(doc, "3. Tecnologias Utilizadas", 1)

# Tabla de tecnologías
tech_table = doc.add_table(rows=1, cols=3)
tech_table.style = 'Table Grid'
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Encabezados
headers = ["Tecnologia", "Uso en el Proyecto", "Version"]
for i, h in enumerate(headers):
    cell = tech_table.rows[0].cells[i]
    set_cell_bg(cell, "023E8A")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)

techs = [
    ("HTML5 Canvas",  "Renderizado del juego 2D",          "HTML 5"),
    ("JavaScript",    "Logica del juego, fisicas, IA",     "ES6+"),
    ("CSS3",          "Interfaz, menus, animaciones UI",   "CSS 3"),
    ("Node.js",       "Servidor local de desarrollo",      "v18+"),
    ("Express.js",    "Servidor HTTP para archivos",       "v4"),
]

for tech, uso, ver in techs:
    row = tech_table.add_row()
    data = [tech, uso, ver]
    for i, txt in enumerate(data):
        cell = row.cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.font.size = Pt(10)
        if i == 0:
            r.bold = True

doc.add_paragraph()
heading(doc, "4. Mecanicas Implementadas", 1)
mecanicas = [
    ("Movimiento del jugador",    "El personaje avanza automaticamente. El jugador controla salto y agacharse."),
    ("Sistema de salto",         "Fisica de gravedad: al presionar W o Espacio se aplica velocidad vertical negativa (jumpForce). La gravedad devuelve el personaje al suelo cada frame."),
    ("Deteccion de colisiones",  "AABB (Axis-Aligned Bounding Box): compara rectangulos del jugador y los obstaculos para detectar impactos."),
    ("Obstaculos",               "Tres tipos: Bloque (esquivar saltando), Puas (peligro en el suelo), Drone (volar, esquivar agachandose)."),
    ("Monedas",                  "Aparecen en patrones de linea o arco. Al recogerlas suman +10 puntos."),
    ("Sistema de puntuacion",    "Puntos por distancia recorrida + bonificacion por monedas. Recordados en localStorage."),
    ("Niveles / Dificultad",     "3 niveles demo con velocidad y frecuencia de obstaculos crecientes."),
    ("Multiplayer local",        "2 jugadores en el mismo teclado: P1 usa W/S, P2 usa Flechas."),
]

for titulo, desc in mecanicas:
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_t = p.add_run(titulo + ": ")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_d = p.add_run(desc)
    run_d.font.size = Pt(11)

doc.add_paragraph()
heading(doc, "5. Arquitectura del Codigo", 1)
body(doc, "El proyecto sigue una arquitectura modular donde cada sistema es una clase independiente:")

arch_table = doc.add_table(rows=1, cols=2)
arch_table.style = 'Table Grid'
arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(["Archivo", "Responsabilidad"]):
    cell = arch_table.rows[0].cells[i]
    set_cell_bg(cell, "023E8A")
    p = cell.paragraphs[0]
    run = p.add_run(h)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)

archivos = [
    ("game.js",      "Bucle principal (game loop), coordina todos los sistemas"),
    ("player.js",    "Clase Player: posicion, fisica, salto, vidas, animacion"),
    ("obstacles.js", "Clase ObstacleSystem: generacion y dibujo de obstaculos"),
    ("coins.js",     "Clase CoinSystem: patrones de monedas y recoleccion"),
    ("collision.js", "Clase CollisionSystem: deteccion AABB de colisiones"),
    ("camera.js",    "Clase Camera: scroll del mundo, zoom, shake de pantalla"),
    ("renderer.js",  "Clase Renderer: fondo, suelo y utilidades de dibujo"),
    ("levels.js",    "Datos de los 3 niveles: velocidad, gravedad, meta"),
    ("score.js",     "Sistema de puntuacion y almacenamiento de records"),
    ("ui.js",        "Manejo de pantallas, menus, HUD y transiciones"),
    ("audio.js",     "Efectos de sonido generados con Web Audio API"),
]

for archivo, resp in archivos:
    row = arch_table.add_row()
    c0 = row.cells[0]
    r0 = c0.paragraphs[0].add_run(archivo)
    r0.bold = True
    r0.font.size = Pt(10)
    r0.font.color.rgb = AZUL_MED

    c1 = row.cells[1]
    r1 = c1.paragraphs[0].add_run(resp)
    r1.font.size = Pt(10)

doc.add_paragraph()
heading(doc, "6. Controles del Juego", 1)

ctrl_table = doc.add_table(rows=1, cols=3)
ctrl_table.style = 'Table Grid'
ctrl_table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, h in enumerate(["Jugador", "Tecla", "Accion"]):
    cell = ctrl_table.rows[0].cells[i]
    set_cell_bg(cell, "023E8A")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.size = Pt(10)

controles = [
    ("Player 1", "W  /  Espacio", "Saltar"),
    ("Player 1", "S",             "Agacharse"),
    ("Player 1", "D  /  Shift Izq", "Dash (impulso)"),
    ("Player 2", "Flecha Arriba", "Saltar"),
    ("Player 2", "Flecha Abajo",  "Agacharse"),
    ("Player 2", "Flecha Der",    "Dash (impulso)"),
    ("General",  "P  /  Esc",    "Pausar el juego"),
]

for jug, tecla, accion in controles:
    row = ctrl_table.add_row()
    vals = [jug, tecla, accion]
    for i, v in enumerate(vals):
        cell = row.cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(v)
        r.font.size = Pt(10)
        if i == 0 and jug != "General":
            r.bold = True

doc.add_paragraph()
heading(doc, "7. Estado Actual y Proximos Pasos", 1)

heading(doc, "Implementado (Avance 1):", 2, color=RGBColor(0x02, 0x8A, 0x0F))
impl = [
    "Motor de fisicas basico (gravedad, salto, colision con el suelo)",
    "Sistema de colisiones AABB",
    "3 tipos de obstaculos (bloque, puas, drone)",
    "Recoleccion de monedas con patrones",
    "Sistema de puntuacion y records en localStorage",
    "3 niveles demo con dificultad creciente",
    "Modo multijugador local en el mismo teclado",
    "HUD (vidas, puntos, monedas, tiempo, progreso)",
    "Menus, pausa, game over y pantalla de resultados",
]
for item in impl:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(item)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x02, 0x8A, 0x0F)

heading(doc, "Proximas Mejoras (Avance 2):", 2, color=RGBColor(0xCC, 0x55, 0x00))
prox = [
    "Sprites / animaciones de personaje con spritesheets",
    "Mas niveles con mapas diseñados a mano",
    "Power-ups (escudo, super salto, x2 puntos)",
    "Musica de fondo procedural mejorada",
    "Sistema de progresion y desbloqueo de personajes",
]
for item in prox:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(item)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xCC, 0x55, 0x00)

doc.add_paragraph()
heading(doc, "8. Conclusiones", 1)
body(doc,
    "Este prototipo demuestra la viabilidad tecnica del proyecto Runner Challenge. "
    "Se implementaron exitosamente las mecanicas fundamentales de un videojuego runner: "
    "fisicas de salto con gravedad, deteccion de colisiones, sistema de puntuacion "
    "y multijugador local. El codigo esta organizado de forma modular, facilitando "
    "la incorporacion de nuevas funcionalidades en los siguientes avances del proyecto."
)

doc.add_paragraph()
body(doc,
    "El mayor reto tecnico fue el sistema de camara compartida para el modo multijugador, "
    "que debio seguir a ambos jugadores simultaneamente aplicando zoom dinamico segun "
    "la separacion entre ellos. Este problema se resolvio con interpolacion suave (lerp) "
    "y un sistema de advertencia visual cuando un jugador va muy atras."
)

# Guardar
doc.save(OUTPUT)
print(f"DOCUMENTO GENERADO: {OUTPUT}")
